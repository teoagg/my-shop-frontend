from pathlib import Path

from docx import Document


DOCX_PATH = Path(r"C:/Users/Thodoris/Downloads/αναφορά_4 (1).docx")

KEYWORDS = [
    "αρχιτεκτον",
    "headless",
    "strapi",
    "next",
    "serverless",
    "mach",
    "stripe",
    "recommender",
    "recommend",
    "a/b",
    "evaluation",
    "αξιολόγ",
    "συμπερά",
    "υλοποίηση",
]


def main():
    doc = Document(DOCX_PATH)
    print(f"FILE: {DOCX_PATH}")
    print(f"PARAGRAPHS: {len(doc.paragraphs)}")
    print("\n## HEADINGS")
    for i, paragraph in enumerate(doc.paragraphs):
      text = paragraph.text.strip()
      style = paragraph.style.name
      if text and ("Heading" in style or "Τίτλος" in style or "Επικεφαλίδα" in style):
          print(f"{i}: {style!r}: {text[:220]}")

    print("\n## KEYWORD HITS")
    for i, paragraph in enumerate(doc.paragraphs):
      text = paragraph.text.strip()
      if not text:
          continue
      lowered = text.lower()
      hits = [keyword for keyword in KEYWORDS if keyword in lowered]
      if hits:
          print(f"{i}: {paragraph.style.name!r}: hits={hits}: {text[:280]}")


if __name__ == "__main__":
    main()
